from __future__ import annotations

from pathlib import Path
from typing import TypeAlias

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .theme import WordPageLayout, WordTheme

import pandas as pd

from .tables import (
    add_action_plan_table as _add_action_plan_table,
    add_data_table,
)

from .figures import add_figure as _add_figure

PathLike: TypeAlias = str | Path


class WordReport:
    """
    Classe base para criação de relatórios Word.

    Responsabilidades:
    - criar ou abrir um documento;
    - configurar tamanho da página e margens;
    - configurar estilos globais;
    - adicionar títulos, subtítulos e parágrafos;
    - inserir quebras de página;
    - salvar o documento.

    Elementos específicos de clientes ou empresas não devem ser
    definidos diretamente nesta classe.
    """

    def __init__(
            self,
            *,
            theme: WordTheme | None = None,
            layout: WordPageLayout | None = None,
            template_path: PathLike | None = None,
    ) -> None:
        self.theme = theme or WordTheme()
        self.layout = layout or WordPageLayout()

        if template_path is None:
            self._document: DocumentObject = Document()

        else:
            template = Path(template_path)

            if not template.exists():
                raise FileNotFoundError(
                    f"Template Word não encontrado: {template}"
                )

            self._document = Document(template)

        self._configure_page_layout()
        self._configure_styles()

    @property
    def document(self) -> DocumentObject:
        """
        Retorna o objeto Document do python-docx.

        Isso permite que aplicações que utilizam dv_ferm_analytics
        tenham acesso a funcionalidades avançadas do python-docx
        sem que a biblioteca precise encapsular tudo.
        """
        return self._document

    def set_header_fields(
            self,
            *,
            usina: str,
            mes: str,
    ) -> None:
        """
        Substitui os campos variáveis existentes no cabeçalho do template.
        """

        replacements = {
            "{{USINA}}": str(usina).strip(),
            "{{MES}}": str(mes).strip(),
        }

        found: set[str] = set()

        for section in self._document.sections:
            header = section.header

            for text_node in header._element.iter(qn("w:t")):
                if not text_node.text:
                    continue

                for placeholder, value in replacements.items():
                    if placeholder in text_node.text:
                        text_node.text = text_node.text.replace(
                            placeholder,
                            value,
                        )
                        found.add(placeholder)

        missing = set(replacements) - found

        if missing:
            raise ValueError(
                "Campos não encontrados no cabeçalho: "
                + ", ".join(sorted(missing))
            )

    def set_footer(
            self,
            *,
            document_code: str,
            show_page_number: bool = True,
    ) -> None:
        """
        Configura o rodapé com código do documento centralizado
        e número automático da página à direita.
        """

        usable_width_cm = (
                self.layout.width_cm
                - self.layout.left_margin_cm
                - self.layout.right_margin_cm
        )

        for section in self._document.sections:
            footer = section.footer
            paragraph = footer.paragraphs[0]

            paragraph.clear()

            paragraph.paragraph_format.tab_stops.add_tab_stop(
                Cm(usable_width_cm / 2),
                WD_TAB_ALIGNMENT.CENTER,
            )

            paragraph.paragraph_format.tab_stops.add_tab_stop(
                Cm(usable_width_cm),
                WD_TAB_ALIGNMENT.RIGHT,
            )

            # Vai até o centro da página
            paragraph.add_run("\t")

            code_run = paragraph.add_run(document_code)
            code_run.font.name = self.theme.body_font
            code_run.font.size = Pt(8)

            if show_page_number:
                # Vai até a margem direita
                paragraph.add_run("\t")

                page_run = paragraph.add_run()
                page_run.font.name = self.theme.body_font
                page_run.font.size = Pt(8)

                field_begin = OxmlElement("w:fldChar")
                field_begin.set(
                    qn("w:fldCharType"),
                    "begin",
                )

                field_code = OxmlElement("w:instrText")
                field_code.text = " PAGE "

                field_end = OxmlElement("w:fldChar")
                field_end.set(
                    qn("w:fldCharType"),
                    "end",
                )

                page_run._r.append(field_begin)
                page_run._r.append(field_code)
                page_run._r.append(field_end)

    def add_title(
        self,
        text: str,
    ):
        """
        Adiciona o título principal do documento.
        """
        paragraph = self._document.add_paragraph(
            style="Title",
        )

        paragraph.add_run(text)

        return paragraph


    def add_heading(
        self,
        text: str,
        *,
        level: int = 1,
    ):
        """
        Adiciona um título de seção.

        Atualmente são suportados os níveis 1 e 2.
        """
        if level not in (1, 2):
            raise ValueError(
                "O nível do título deve ser 1 ou 2."
            )

        return self._document.add_heading(
            text,
            level=level,
        )

    def add_paragraph(
        self,
        text: str = "",
        *,
        align: str | None = None,
    ):
        """
        Adiciona um parágrafo de texto.

        align pode ser:
        - None
        - 'left'
        - 'center'
        - 'right'
        - 'justify'
        """
        paragraph = self._document.add_paragraph(
            text,
            style="Normal",
        )

        if align is not None:
            paragraph.alignment = self._get_alignment(
                align
            )

        return paragraph

    def add_data_table(
            self,
            df: pd.DataFrame,
            *,
            include_index: bool = False,
            higher_is_better: set[str] | None = None,
            lower_is_better: set[str] | None = None,
    ):
        return add_data_table(
            self._document,
            df,
            include_index=include_index,
            higher_is_better=higher_is_better,
            lower_is_better=lower_is_better,
        )

    def add_data_table(
            self,
            df: pd.DataFrame,
            *,
            include_index: bool = False,
            higher_is_better: set[str] | None = None,
            lower_is_better: set[str] | None = None,
            column_widths_cm: list[float] | None = None,
            decimal_places: int = 2,
    ):
        return add_data_table(
            self._document,
            df,
            include_index=include_index,
            higher_is_better=higher_is_better,
            lower_is_better=lower_is_better,
            column_widths_cm=column_widths_cm,
            decimal_places=decimal_places,
        )

    def add_page_break(self) -> None:
        """
        Insere uma quebra de página.
        """
        self._document.add_page_break()

    def save(
        self,
        output_path: PathLike,
    ) -> Path:
        """
        Salva o documento Word.

        A pasta de destino é criada automaticamente,
        caso ainda não exista.
        """
        path = Path(output_path)

        if path.suffix.lower() != ".docx":
            raise ValueError(
                "O arquivo de saída deve possuir extensão .docx."
            )

        path.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        self._document.save(path)

        return path

    def _configure_page_layout(self) -> None:
        """
        Configura tamanho físico da página,
        margens, cabeçalho e rodapé.
        """
        for section in self._document.sections:
            section.page_width = Cm(
                self.layout.width_cm
            )

            section.page_height = Cm(
                self.layout.height_cm
            )

            section.top_margin = Cm(
                self.layout.top_margin_cm
            )

            section.bottom_margin = Cm(
                self.layout.bottom_margin_cm
            )

            section.left_margin = Cm(
                self.layout.left_margin_cm
            )

            section.right_margin = Cm(
                self.layout.right_margin_cm
            )

            section.header_distance = Cm(
                self.layout.header_distance_cm
            )

            section.footer_distance = Cm(
                self.layout.footer_distance_cm
            )

    def _configure_styles(self) -> None:
        """
        Configura os estilos básicos utilizados
        pelo documento.
        """
        styles = self._document.styles

        # -------------------------------------------------
        # Texto normal
        # -------------------------------------------------
        normal = styles["Normal"]

        self._set_style_font(
            normal,
            font_name=self.theme.body_font,
            size_pt=self.theme.body_size_pt,
            color_hex=self.theme.body_color,
        )

        normal.paragraph_format.space_after = Pt(8)
        normal.paragraph_format.line_spacing = 1.08

        # -------------------------------------------------
        # Título principal
        # -------------------------------------------------
        title = styles["Title"]

        self._set_style_font(
            title,
            font_name=self.theme.title_font,
            size_pt=self.theme.title_size_pt,
            color_hex=self.theme.title_color,
        )

        title.paragraph_format.space_after = Pt(12)
        title.paragraph_format.line_spacing = 1.0

        # -------------------------------------------------
        # Heading 1
        # -------------------------------------------------
        heading_1 = styles["Heading 1"]

        self._set_style_font(
            heading_1,
            font_name=self.theme.heading_font,
            size_pt=self.theme.heading_1_size_pt,
            color_hex=self.theme.title_color,
        )

        heading_1.paragraph_format.space_before = Pt(24)
        heading_1.paragraph_format.space_after = Pt(6)

        # -------------------------------------------------
        # Heading 2
        # -------------------------------------------------
        heading_2 = styles["Heading 2"]

        self._set_style_font(
            heading_2,
            font_name=self.theme.heading_font,
            size_pt=self.theme.heading_2_size_pt,
            color_hex=self.theme.accent_color,
        )

        heading_2.paragraph_format.space_before = Pt(12)
        heading_2.paragraph_format.space_after = Pt(4)

    @staticmethod
    def _set_style_font(
        style,
        *,
        font_name: str,
        size_pt: float,
        color_hex: str,
    ) -> None:
        """
        Aplica fonte, tamanho e cor ao estilo Word.
        """
        style.font.name = font_name
        style.font.size = Pt(size_pt)

        style.font.color.rgb = RGBColor.from_string(
            color_hex
        )

        r_pr = style.element.get_or_add_rPr()

        r_pr.rFonts.set(
            qn("w:ascii"),
            font_name,
        )

        r_pr.rFonts.set(
            qn("w:hAnsi"),
            font_name,
        )

        r_pr.rFonts.set(
            qn("w:eastAsia"),
            font_name,
        )

    @staticmethod
    def _get_alignment(
        alignment: str,
    ):
        """
        Converte uma string em alinhamento do Word.
        """
        options = {
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
            "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        }

        key = alignment.lower().strip()

        if key not in options:
            raise ValueError(
                "align deve ser 'left', 'center', "
                "'right' ou 'justify'."
            )

        return options[key]

    def add_figure(
            self,
            figure,
            *,
            width_cm: float | None = 16.0,
            height_cm: float | None = None,
            alignment: str = "center",
            dpi: int = 200,
            caption: str | None = None,
            caption_font_size_pt: float = 8.0,
            space_after_pt: float = 6.0,
    ):
        """
        Adiciona uma figura ou imagem ao relatório.
        """
        return _add_figure(
            self._document,
            figure,
            width_cm=width_cm,
            height_cm=height_cm,
            alignment=alignment,
            dpi=dpi,
            caption=caption,
            caption_font_size_pt=caption_font_size_pt,
            space_after_pt=space_after_pt,
        )


