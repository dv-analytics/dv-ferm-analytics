from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import BinaryIO

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from matplotlib.figure import Figure


AlignmentValue = str | WD_ALIGN_PARAGRAPH


def _validate_positive(
    value: float | None,
    *,
    name: str,
) -> None:
    """
    Valida valores positivos opcionais.
    """

    if value is None:
        return

    if value <= 0:
        raise ValueError(
            f"{name} deve ser maior que zero."
        )


def _validate_non_negative(
    value: float,
    *,
    name: str,
) -> None:
    """
    Valida valores que podem ser zero, mas não negativos.
    """

    if value < 0:
        raise ValueError(
            f"{name} não pode ser negativo."
        )


def _resolve_alignment(
    alignment: AlignmentValue,
):
    """
    Converte alinhamento textual para WD_ALIGN_PARAGRAPH.
    """

    if not isinstance(
        alignment,
        str,
    ):
        return alignment

    normalized = (
        alignment
        .strip()
        .lower()
    )

    mapping = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }

    if normalized not in mapping:
        raise ValueError(
            "alignment deve ser um dos valores: "
            "'left', 'center', 'right' ou 'justify'."
        )

    return mapping[
        normalized
    ]


def _matplotlib_figure_to_stream(
    figure: Figure,
    *,
    dpi: int,
) -> BytesIO:
    """
    Converte uma figura Matplotlib em PNG em memória.
    """

    stream = BytesIO()

    figure.savefig(
        stream,
        format="png",
        dpi=dpi,
        bbox_inches="tight",
    )

    stream.seek(
        0
    )

    return stream


def _prepare_image_source(
    figure: Figure | str | Path | BinaryIO | BytesIO,
    *,
    dpi: int,
):
    """
    Normaliza a origem da imagem para o python-docx.

    O python-docx aceita:
    - caminho como string;
    - stream binário seekable.

    Pathlib.Path é convertido explicitamente para string.
    Isso evita o erro:

        AttributeError:
        'WindowsPath' object has no attribute 'seek'
    """

    if isinstance(
        figure,
        Figure,
    ):
        return _matplotlib_figure_to_stream(
            figure,
            dpi=dpi,
        )

    if isinstance(
        figure,
        (
            str,
            Path,
        ),
    ):
        image_path = Path(
            figure
        )

        if not image_path.exists():
            raise FileNotFoundError(
                "Imagem não encontrada:\n"
                f"{image_path}"
            )

        if not image_path.is_file():
            raise ValueError(
                "O caminho informado não é um arquivo:\n"
                f"{image_path}"
            )

        # IMPORTANTE:
        # retornar str e não Path.
        return str(
            image_path
        )

    if (
        hasattr(
            figure,
            "read",
        )
        and hasattr(
            figure,
            "seek",
        )
    ):
        figure.seek(
            0
        )

        return figure

    # Compatibilidade com objetos figure-like.
    if hasattr(
        figure,
        "savefig",
    ):
        stream = BytesIO()

        figure.savefig(
            stream,
            format="png",
            dpi=dpi,
            bbox_inches="tight",
        )

        stream.seek(
            0
        )

        return stream

    raise TypeError(
        "figure deve ser uma figura Matplotlib, "
        "um caminho str/Path ou um stream binário."
    )


def add_figure(
    document,
    figure: Figure | str | Path | BinaryIO | BytesIO,
    *,
    width_cm: float | None = 16.0,
    height_cm: float | None = None,
    alignment: AlignmentValue = "center",
    dpi: int = 200,
    caption: str | None = None,
    caption_font_size_pt: float = 9.0,
    space_after_pt: float = 6.0,
):
    """
    Adiciona uma figura ao documento Word.

    Esta assinatura mantém compatibilidade com WordReport.add_figure().

    Parameters
    ----------
    document
        Documento python-docx.
    figure
        Figura Matplotlib, caminho str/Path ou stream binário.
    width_cm
        Largura em centímetros.
    height_cm
        Altura em centímetros. Se None, preserva proporção.
    alignment
        left, center, right, justify ou WD_ALIGN_PARAGRAPH.
    dpi
        DPI usado ao converter figuras Matplotlib.
    caption
        Legenda opcional.
    caption_font_size_pt
        Tamanho da fonte da legenda em pontos.
    space_after_pt
        Espaço após figura e legenda, em pontos.

    Returns
    -------
    InlineShape
        Objeto retornado por python-docx.
    """

    _validate_positive(
        width_cm,
        name="width_cm",
    )

    _validate_positive(
        height_cm,
        name="height_cm",
    )

    _validate_positive(
        float(dpi),
        name="dpi",
    )

    _validate_positive(
        caption_font_size_pt,
        name="caption_font_size_pt",
    )

    _validate_non_negative(
        space_after_pt,
        name="space_after_pt",
    )

    paragraph_alignment = (
        _resolve_alignment(
            alignment
        )
    )

    image_source = (
        _prepare_image_source(
            figure,
            dpi=dpi,
        )
    )

    paragraph = (
        document.add_paragraph()
    )

    paragraph.alignment = (
        paragraph_alignment
    )

    paragraph.paragraph_format.space_after = (
        Pt(
            space_after_pt
        )
    )

    run = (
        paragraph.add_run()
    )

    picture_kwargs = {}

    if width_cm is not None:
        picture_kwargs[
            "width"
        ] = Cm(
            width_cm
        )

    if height_cm is not None:
        picture_kwargs[
            "height"
        ] = Cm(
            height_cm
        )

    picture = (
        run.add_picture(
            image_source,
            **picture_kwargs,
        )
    )

    if caption:
        caption_paragraph = (
            document.add_paragraph()
        )

        caption_paragraph.alignment = (
            paragraph_alignment
        )

        caption_paragraph.paragraph_format.space_after = (
            Pt(
                space_after_pt
            )
        )

        caption_run = (
            caption_paragraph.add_run(
                str(
                    caption
                )
            )
        )

        caption_run.italic = True

        caption_run.font.size = (
            Pt(
                caption_font_size_pt
            )
        )

    return picture
