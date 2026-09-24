from io import BytesIO
from pathlib import Path

import matplotlib.pyplot as plt
from docx import Document

from dv_ferm_analytics.word.figures import (
    add_figure,
)


def _create_png(
    image_path: Path,
) -> None:
    fig, ax = plt.subplots()

    ax.plot(
        [0, 1, 2],
        [1, 3, 2],
    )

    fig.savefig(
        image_path,
        dpi=100,
    )

    plt.close(
        fig
    )


def test_add_figure_accepts_pathlib_path(
    tmp_path: Path,
):
    image_path = (
        tmp_path
        / "figure.png"
    )

    _create_png(
        image_path
    )

    document = Document()

    picture = add_figure(
        document,
        image_path,
        width_cm=8,
    )

    assert picture is not None

    output_path = (
        tmp_path
        / "report.docx"
    )

    document.save(
        output_path
    )

    assert output_path.exists()
    assert output_path.stat().st_size > 0


def test_add_figure_accepts_string_path(
    tmp_path: Path,
):
    image_path = (
        tmp_path
        / "figure.png"
    )

    _create_png(
        image_path
    )

    document = Document()

    picture = add_figure(
        document,
        str(
            image_path
        ),
        width_cm=8,
    )

    assert picture is not None


def test_add_figure_accepts_matplotlib_figure():
    fig, ax = plt.subplots()

    ax.plot(
        [0, 1],
        [1, 0],
    )

    document = Document()

    picture = add_figure(
        document,
        fig,
        width_cm=8,
    )

    plt.close(
        fig
    )

    assert picture is not None


def test_add_figure_accepts_binary_stream():
    fig, ax = plt.subplots()

    ax.plot(
        [0, 1],
        [1, 2],
    )

    stream = BytesIO()

    fig.savefig(
        stream,
        format="png",
        dpi=100,
    )

    plt.close(
        fig
    )

    stream.seek(
        0
    )

    document = Document()

    picture = add_figure(
        document,
        stream,
        width_cm=8,
    )

    assert picture is not None


def test_add_figure_accepts_caption_font_size(
    tmp_path: Path,
):
    image_path = (
        tmp_path
        / "figure.png"
    )

    _create_png(
        image_path
    )

    document = Document()

    picture = add_figure(
        document,
        image_path,
        width_cm=8,
        caption="Figura de teste",
        caption_font_size_pt=8.5,
        space_after_pt=4,
    )

    assert picture is not None

    assert (
        document.paragraphs[
            -1
        ].text
        == "Figura de teste"
    )


def test_add_figure_preserves_wordreport_signature(
    tmp_path: Path,
):
    """
    Reproduz os argumentos enviados por WordReport.add_figure().
    """

    image_path = (
        tmp_path
        / "figure.png"
    )

    _create_png(
        image_path
    )

    document = Document()

    picture = add_figure(
        document,
        str(
            image_path
        ),
        width_cm=16.2,
        height_cm=None,
        alignment="center",
        dpi=200,
        caption=None,
        caption_font_size_pt=9.0,
        space_after_pt=6.0,
    )

    assert picture is not None
